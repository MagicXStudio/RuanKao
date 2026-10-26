# -*- coding: utf-8 -*-
"""docx -> markdown 转换器（历年真题专用）"""
import re, sys, os, zipfile
from pathlib import Path
import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC_DIR = Path(r"F:/RuanKao/系统集成项目管理工程师")
OUT_DIR = Path(r"F:/RuanKao/历年真题MD")

JOBS = [
    ("2020年5月系统集成项目管理工程师考试葵花宝典之金色考点暨历年真题解析_1-410.docx",
     "01-历年真题解析_第1-410页.md"),
    ("2020年5月系统集成项目管理工程师考试葵花宝典之金色考点暨历年真题解析_411-820.docx",
     "02-历年真题解析_第411-820页.md"),
    ("2020年5月系统集成项目管理工程师考试葵花宝典之金色考点暨历年真题解析_821-1224.docx",
     "03-历年真题解析_第821-1224页.md"),
    ("2020年5月系统集成项目管理工程师考试葵花宝典之金色题库.docx",
     "04-历年真题题库_按章分类.md"),
    ("历年案例分析-解决措施参考.docx",
     "05-历年案例分析_解决措施参考.md"),
]

RE_PAGE_NO = re.compile(r'^\s*[-–—\s]*\d{1,4}\s*[-–—\s]*$')
RE_CHAPTER = re.compile(r'^第[一二三四五六七八九十百\d]+章')
RE_TOPIC = re.compile(r'^重要考点\s*\d+')
RE_NUM_SUB = re.compile(r'^\d{1,3}、[^\d（【，。；]{2,28}$')
RE_BRACKET = re.compile(r'^【[^】]{1,20}】')

IMG_DIR = OUT_DIR / "images"


class ImageExtractor:
    def __init__(self, doc, slug):
        self.doc = doc
        self.slug = slug
        self.dir = IMG_DIR / slug
        self.dir.mkdir(parents=True, exist_ok=True)
        self.counter = 0

    def harvest(self, paragraph):
        """提取段落中所有图片，返回相对 md 路径列表"""
        out = []
        # 现代 drawing
        for blip in paragraph._p.findall('.//' + qn('a:blip')):
            rid = blip.get(qn('r:embed'))
            if not rid:
                continue
            try:
                part = self.doc.part.related_parts[rid]
            except KeyError:
                continue
            out.append(self._save(part))
        # 旧式 w:pict / v:imagedata（VML 命名空间不在 python-docx nsmap 中，写死 URI）
        v_imgdata = '{urn:schemas-microsoft-com:vml}imagedata'
        for imgdata in paragraph._p.findall('.//' + v_imgdata):
            rid = imgdata.get(qn('r:id'))
            if not rid:
                continue
            try:
                part = self.doc.part.related_parts[rid]
            except KeyError:
                continue
            out.append(self._save(part))
        return out

    def _save(self, part):
        self.counter += 1
        ext = os.path.splitext(str(part.partname))[1] or ".png"
        name = f"{self.slug}_{self.counter:04d}{ext}"
        (self.dir / name).write_bytes(part.blob)
        return f"images/{self.slug}/{name}"


def cell_text(cell):
    parts = []
    for p in cell.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    return "<br>".join(parts).replace("|", "\\|")


def table_md(tbl):
    rows = tbl.rows
    if not rows:
        return ""
    lines = []
    ncol = max(len(r.cells) for r in rows)
    for i, r in enumerate(rows):
        cells = [cell_text(c) for c in r.cells]
        cells += [""] * (ncol - len(cells))
        lines.append("| " + " | ".join(cells) + " |")
        if i == 0:
            lines.append("|" + "---|" * ncol)
    return "\n".join(lines)


def para_md(p, imgx, img_in_table=False):
    text = p.text.strip()
    imgs = imgx.harvest(p)
    if imgs and not text:
        return "\n".join(f"![图]({u})" for u in imgs)
    if not text:
        return ""
    if RE_PAGE_NO.match(text):
        return ""
    if RE_CHAPTER.match(text):
        return f"## {text}"
    if RE_TOPIC.match(text):
        return f"### {text}"
    if RE_BRACKET.match(text):
        return f"#### {text}"
    if not img_in_table and RE_NUM_SUB.match(text):
        return f"### {text}"
    if imgs:
        return text + "\n" + "\n".join(f"![图]({u})" for u in imgs)
    return text


def convert(docx_path, md_path):
    docx_path, md_path = Path(docx_path), Path(md_path)
    doc = docx.Document(docx_path)
    slug = md_path.stem.replace("-", "_")
    imgx = ImageExtractor(doc, slug)
    lines = [f"# {md_path.stem}", ""]

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            md = para_md(Paragraph(child, doc), imgx)
            if md:
                lines.append(md)
                lines.append("")
        elif child.tag == qn('w:tbl'):
            md = table_md(Table(child, doc))
            if md:
                lines.append(md)
                lines.append("")

    text = "\n".join(lines)
    text = re.sub(r'\n{3,}', '\n\n', text)
    OUT_DIR.joinpath(md_path).write_text(text, encoding="utf-8")
    n_img = imgx.counter
    n_para = len(doc.paragraphs)
    n_tbl = len(doc.tables)
    print(f"OK  {md_path}  段落:{n_para} 表格:{n_tbl} 图片:{n_img} 大小:{len(text)//1024}KB")


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for src, dst in JOBS:
        convert(SRC_DIR / src, dst)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
